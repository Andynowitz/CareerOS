from __future__ import annotations

import io

from docx import Document
from pypdf import PdfReader


class ResumeParserError(Exception):
    """Raised when resume text extraction fails."""


def parse_resume(data: bytes, content_type: str) -> str:
    try:
        if content_type == "application/pdf":
            reader = PdfReader(io.BytesIO(data))
            return "\n".join(
                page.extract_text() or ""
                for page in reader.pages
            ).strip()

        if content_type == (
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        ):
            document = Document(io.BytesIO(data))

            parts = [
                paragraph.text
                for paragraph in document.paragraphs
                if paragraph.text.strip()
            ]

            for table in document.tables:
                for row in table.rows:
                    parts.append(
                        " | ".join(cell.text.strip() for cell in row.cells)
                    )

            return "\n".join(parts).strip()

    except Exception as exc:
        raise ResumeParserError(
            "Could not extract text from the resume"
        ) from exc

    raise ResumeParserError("Unsupported resume format")